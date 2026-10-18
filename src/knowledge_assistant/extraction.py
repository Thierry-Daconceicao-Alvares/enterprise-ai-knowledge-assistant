"""Validate uploaded files and extract their text, page by page.

Supported formats:

- PDF: text of each page, with its page number (PDFium engine). Scanned PDFs
  have no text layer and are rejected: OCR is out of scope.
- DOCX: paragraphs, lists and tables in document order. Word documents have no
  fixed pages (pagination depends on the viewer), so citations point to the
  section instead of a page.
- TXT: UTF-8 text (Windows-1252 accepted as a fallback), no pages either.

The extracted text is cleaned without changing its meaning: whitespace and
control characters are normalised, and running headers, footers and page
numbers are removed from PDF pages.
"""

from __future__ import annotations

import hashlib
import io
import itertools
import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import PurePath

from .errors import (
    EmptyFileError,
    FileTooLargeError,
    MalformedDocumentError,
    NoExtractableTextError,
    UnsupportedFileTypeError,
)

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt"}
# Limits on the content of a DOCX file (a ZIP archive), to refuse "zip bombs"
# that are small on disk but huge once decompressed.
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2000
MIN_TEXT_CHARACTERS = 20
SUPERSEDED_STATUSES = frozenset({"superseded", "obsolete", "withdrawn", "archived", "retired"})


@dataclass(frozen=True)
class Page:
    """Text of one page. `number` is None for formats without fixed pages."""

    number: int | None
    text: str


@dataclass(frozen=True)
class DocumentMetadata:
    """Metadata read from the document itself. Only the title is always set."""

    title: str
    doc_id: str | None = None
    version: str | None = None
    effective_date: str | None = None  # ISO 8601
    status: str | None = None  # as declared in the document, lower case
    owner: str | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    file_name: str
    file_type: str
    sha256: str
    size_bytes: int
    pages: list[Page]
    metadata: DocumentMetadata
    page_count: int | None
    warnings: list[str] = field(default_factory=list)

    @property
    def doc_key(self) -> str:
        """Identifier derived from the file content: the same file uploaded
        twice gets the same key, whatever its name."""
        return self.sha256[:16]

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text)


# ---------------------------------------------------------------------- validation


_UNSAFE_NAME_CHARACTERS = re.compile(r"[^\w.\- ()]+", re.UNICODE)


def safe_file_name(name: str) -> str:
    """The base name of an uploaded file, without any directory part and with
    unusual characters replaced, so it can be stored and displayed safely."""
    base = PurePath(name.replace("\\", "/")).name.strip()
    base = _UNSAFE_NAME_CHARACTERS.sub("_", unicodedata.normalize("NFC", base))
    base = base.lstrip(".") or "document"
    if len(base) > 120:
        stem, dot, extension = base.rpartition(".")
        base = (stem[: 115 - len(extension)] + dot + extension) if dot else base[:120]
    return base


def detect_file_type(file_name: str) -> str:
    extension = PurePath(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        shown = extension or "files without an extension"
        raise UnsupportedFileTypeError(
            f"{file_name}: {shown} is not supported. Upload PDF, DOCX or TXT files."
        )
    return SUPPORTED_EXTENSIONS[extension]


def validate_upload(file_name: str, data: bytes, max_bytes: int) -> str:
    """Check the extension, the size and the file signature before any parsing.
    Returns the file type ('pdf', 'docx' or 'txt')."""
    file_type = detect_file_type(file_name)
    if not data:
        raise EmptyFileError(f"{file_name} is empty.")
    if len(data) > max_bytes:
        raise FileTooLargeError(
            f"{file_name} is {len(data) / 1024 / 1024:.1f} MB; the limit is "
            f"{max_bytes / 1024 / 1024:.0f} MB."
        )
    if file_type == "pdf" and b"%PDF-" not in data[:1024]:
        raise MalformedDocumentError(f"{file_name} is not a valid PDF file.")
    if file_type == "docx":
        _check_docx_archive(file_name, data)
    if file_type == "txt" and b"\x00" in data[:8192]:
        raise MalformedDocumentError(f"{file_name} looks like a binary file, not a text file.")
    return file_type


def _check_docx_archive(file_name: str, data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            uncompressed = sum(entry.file_size for entry in entries)
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise MalformedDocumentError(f"{file_name} is not a valid DOCX file.") from exc
    if "word/document.xml" not in names:
        raise MalformedDocumentError(f"{file_name} is not a Word document.")
    if len(entries) > MAX_DOCX_ENTRIES or uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
        raise FileTooLargeError(f"{file_name} is too large once decompressed.")


# ---------------------------------------------------------------------- extraction


def extract_document(
    file_name: str, data: bytes, max_bytes: int = 10 * 1024 * 1024, max_pages: int = 300
) -> ExtractedDocument:
    """Validate a file and extract its text and metadata.

    Raises a DocumentError subclass with a clear message when the file cannot
    be used (unsupported, empty, too large, corrupted or without text).
    """
    file_type = validate_upload(file_name, data, max_bytes)
    warnings: list[str] = []
    if file_type == "pdf":
        raw_pages, embedded_title = _extract_pdf(file_name, data, max_pages)
        pages = [Page(number, text) for number, text in raw_pages]
    elif file_type == "docx":
        text, embedded_title = _extract_docx(file_name, data)
        pages = [Page(None, text)]
    else:
        text, encoding_warning = _decode_text(data)
        if encoding_warning:
            warnings.append(encoding_warning)
        pages, embedded_title = [Page(None, text)], None

    pages = [Page(page.number, clean_text(page.text)) for page in pages]
    if file_type == "pdf":
        pages = remove_page_furniture(pages, embedded_title)
        empty = [str(page.number) for page in pages if not page.text]
        if empty and len(empty) < len(pages):
            warnings.append(f"No text on page(s) {', '.join(empty)} (images are not read).")

    full_text = "\n".join(page.text for page in pages)
    if sum(char.isalnum() for char in full_text) < MIN_TEXT_CHARACTERS:
        hint = " It may be a scanned document: OCR is not supported." if file_type == "pdf" else ""
        raise NoExtractableTextError(f"{file_name} contains no extractable text.{hint}")

    return ExtractedDocument(
        file_name=file_name,
        file_type=file_type,
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        pages=pages,
        metadata=parse_metadata(full_text, embedded_title, file_name),
        page_count=len(pages) if file_type == "pdf" else None,
        warnings=warnings,
    )


def _extract_pdf(
    file_name: str, data: bytes, max_pages: int
) -> tuple[list[tuple[int, str]], str | None]:
    import pypdfium2 as pdfium  # imported here: only needed for PDF files

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        reason = (
            "it is protected by a password" if "password" in str(exc).lower() else "it is damaged"
        )
        raise MalformedDocumentError(f"{file_name} cannot be read: {reason}.") from exc
    try:
        if len(pdf) == 0:
            raise NoExtractableTextError(f"{file_name} has no pages.")
        if len(pdf) > max_pages:
            raise FileTooLargeError(
                f"{file_name} has {len(pdf)} pages; the limit is {max_pages} pages."
            )
        pages = []
        for index in range(len(pdf)):
            page = pdf[index]
            text_page = page.get_textpage()
            pages.append((index + 1, text_page.get_text_range()))
            text_page.close()
            page.close()
        title = (pdf.get_metadata_dict().get("Title") or "").strip() or None
        return pages, title
    except pdfium.PdfiumError as exc:
        raise MalformedDocumentError(f"{file_name} cannot be read: it is damaged.") from exc
    finally:
        pdf.close()


def _extract_docx(file_name: str, data: bytes) -> tuple[str, str | None]:
    """Paragraphs and tables in document order. Headings are marked with `#`
    so that the chunker can find the sections, and list items keep a marker."""
    from docx import Document  # imported here: only needed for DOCX files
    from docx.table import Table

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises many exception types
        raise MalformedDocumentError(f"{file_name} cannot be read: it is damaged.") from exc

    lines: list[str] = []
    title: str | None = (document.core_properties.title or "").strip() or None
    list_number = 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            list_number = 0
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    text = " ".join(cell.text.split())
                    if text and (not cells or cells[-1] != text):  # merged cells repeat
                        cells.append(text)
                if cells:
                    lines.append(" | ".join(cells))
            continue
        text = block.text.strip()
        style = (block.style.name if block.style is not None else "") or ""
        if not text:
            continue
        if style == "Title":
            title = title or text
            lines.append(text)
        elif style.startswith("Heading"):
            level = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
            lines.append("#" * min(level + 1, 6) + " " + text)
        elif style.startswith("List Number"):
            list_number += 1
            lines.append(f"{list_number}) {text}")
            continue
        elif style.startswith("List"):
            lines.append(f"• {text}")
        else:
            lines.append(text)
        list_number = 0
    # Word paragraphs are complete units: a blank line between them prevents
    # the chunker from merging them like wrapped lines of a PDF.
    return "\n\n".join(lines), title


def _decode_text(data: bytes) -> tuple[str, str | None]:
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "Not UTF-8: decoded as Windows-1252."


# ------------------------------------------------------------------------ cleaning

_CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f​­]")
_SPACES = re.compile(r"[ \t   ]+")
_MANY_BLANK_LINES = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    """Normalise whitespace and characters without changing the meaning:
    unify line breaks, remove control characters and soft hyphens, replace
    ligatures (NFKC) and collapse repeated spaces."""
    text = unicodedata.normalize("NFKC", text.replace("\r\n", "\n").replace("\r", "\n"))
    text = _CONTROL.sub("", text)
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    return _MANY_BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


# "Page 3 of 12" or "page 3/12" anywhere in the line, or a line that is only a
# page number: "3", "3/12", "- 3 -" or "Page 3".
_PAGE_NUMBER_LINE = re.compile(
    r"\bpage\s+\d+\s*(of|/)\s*\d+\b|^\s*(page\s+)?[-–]?\s*\d{1,4}\s*(/\s*\d{1,4})?\s*[-–]?\s*$",
    re.I,
)
_DIGITS = re.compile(r"\d+")


def remove_page_furniture(pages: list[Page], title: str | None) -> list[Page]:
    """Remove running headers, footers and page numbers from PDF pages.

    These lines are repeated on every page and would otherwise end up in the
    middle of the passages. The rules are deliberately conservative:
    - a line in the first or last 3 lines of a page that contains a page number;
    - a first line of pages 2+ that repeats the document title;
    - a line among the last 3 lines that is repeated (digits ignored) at the
      bottom of at least half of the pages, for documents of 2 pages or more.
    """
    if not pages:
        return pages
    split = [page.text.split("\n") for page in pages]
    bottom_counts: Counter[str] = Counter()
    for lines in split:
        for line in set(lines[-3:]):
            bottom_counts[_DIGITS.sub("#", line)] += 1
    repeated_bottom = {
        line
        for line, count in bottom_counts.items()
        if line.strip() and len(pages) >= 2 and count >= max(2, len(pages) / 2)
    }
    title_key = (title or "").strip().lower()

    cleaned: list[Page] = []
    for page, lines in zip(pages, split, strict=True):
        keep = []
        last = len(lines) - 1
        for position, line in enumerate(lines):
            near_top, near_bottom = position < 3, position > last - 3
            if (near_top or near_bottom) and line.strip() and _PAGE_NUMBER_LINE.search(line):
                continue
            if near_bottom and _DIGITS.sub("#", line) in repeated_bottom:
                continue
            running_title = position == 0 and (page.number or 0) > 1 and bool(title_key)
            if running_title and line.strip().lower() == title_key:
                continue
            keep.append(line)
        cleaned.append(Page(page.number, "\n".join(keep).strip()))
    return cleaned


# ------------------------------------------------------------------------ metadata

_DOC_ID = re.compile(r"Document ID\s*:\s*([A-Za-z0-9][\w./-]{1,40})", re.I)
_VERSION = re.compile(r"\bVersion\s*:?\s*v?(\d+(?:\.\d+){0,3})\b", re.I)
_EFFECTIVE = re.compile(r"Effective date\s*:\s*([^|\n]{6,30})", re.I)
_STATUS = re.compile(r"\bStatus\s*:\s*([A-Za-z]+)", re.I)
_OWNER = re.compile(r"\bOwner\s*:\s*([^|\n]{2,80})", re.I)
_DATE_FORMATS = ("%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%B %d, %Y", "%d %b %Y")


def _parse_date(raw: str) -> str | None:
    raw = raw.strip().rstrip(".")
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_metadata(text: str, embedded_title: str | None, file_name: str) -> DocumentMetadata:
    """Read the document ID, version, effective date, status and owner from the
    header block of a document, when it has one (only the beginning of the
    text is searched). The title comes from the file properties, or the line
    before "Document ID", or the first line, or the file name."""
    head = text[:3000]
    doc_id = _first(_DOC_ID, head)
    title = embedded_title
    lines = [line.strip() for line in head.split("\n") if line.strip()]
    if not title and doc_id:
        for previous, line in itertools.pairwise(lines):
            if _DOC_ID.search(line) and len(previous) <= 120:
                title = previous
                break
    if not title:
        title = next((line for line in lines if len(line.split()) >= 2 and len(line) <= 120), None)
    if not title:
        title = PurePath(file_name).stem.replace("_", " ").replace("-", " ")
    status = _first(_STATUS, head)
    effective = _first(_EFFECTIVE, head)
    return DocumentMetadata(
        title=title.lstrip("# ").strip(),
        doc_id=doc_id.rstrip(".") if doc_id else None,
        version=_first(_VERSION, head),
        effective_date=_parse_date(effective) if effective else None,
        status=status.lower() if status else None,
        owner=(_first(_OWNER, head) or "").strip() or None,
    )


def _first(pattern: re.Pattern[str], text: str) -> str | None:
    match = pattern.search(text)
    return match.group(1).strip() if match else None


def version_key(version: str | None) -> tuple[int, ...]:
    """'2.10' sorts after '2.9'; a missing or unreadable version sorts first."""
    if not version:
        return (-1,)
    try:
        return tuple(int(part) for part in version.lstrip("vV").split("."))
    except ValueError:
        return (-1,)


def parse_iso_date(value: str | None) -> date | None:
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None
