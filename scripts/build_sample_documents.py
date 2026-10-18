"""Render the synthetic Nexora documents from Markdown sources to PDF, DOCX and TXT.

Usage:
    python scripts/build_sample_documents.py

Each source in data/sample_documents_src/ starts with a small front matter block
(document ID, title, version, status, output format and file name). The output
files in data/sample_documents/ look like ordinary company documents: a title
block, numbered sections, lists, tables, and page headers and footers in the PDFs.

The generated files are committed to the repository, so this script only needs
to run again after a source changes. It needs fpdf2 and python-docx (installed
with `pip install -e ".[dev]"`) and the DejaVu Sans font, which is preinstalled
on most Linux systems (set NEXORA_FONT_DIR to use another folder).
"""

from __future__ import annotations

import os
import re
import sys
import textwrap
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF
from fpdf.fonts import FontFace

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "sample_documents_src"
OUTPUT_DIR = ROOT / "data" / "sample_documents"

DISCLAIMER = (
    "Nexora Industries is a fictional company. This synthetic document was created "
    "for a portfolio demonstration and does not describe a real organisation."
)
EXTERNAL_DISCLAIMER = "Fictional supplier document, created for a portfolio demonstration."
FIXED_TIMESTAMP = datetime(2026, 10, 1, 9, 0, tzinfo=UTC)
# After a multi-line cell, continue at the left margin on the next line.
NEXT_LINE = {"new_x": "LMARGIN", "new_y": "NEXT"}
FONT_DIRS = [
    os.environ.get("NEXORA_FONT_DIR", ""),
    "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/dejavu",
    "/Library/Fonts",
    "C:/Windows/Fonts",
]


# ----------------------------------------------------------------------------- parsing


@dataclass
class Block:
    kind: str  # heading, paragraph, bullets, numbered, table, hidden
    text: str = ""
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class SourceDocument:
    meta: dict[str, str]
    blocks: list[Block]

    @property
    def external(self) -> bool:
        return self.meta.get("external", "").lower() == "true"


def parse_source(path: Path) -> SourceDocument:
    text = path.read_text(encoding="utf-8")
    match = re.match(r"---\n(.*?)\n---\n", text, re.S)
    if not match:
        raise ValueError(f"{path.name}: missing front matter")
    meta = {}
    for line in match.group(1).splitlines():
        key, _, value = line.partition(":")
        meta[key.strip()] = value.strip()
    for required in ("title", "format", "file"):
        if required not in meta:
            raise ValueError(f"{path.name}: front matter needs '{required}'")
    return SourceDocument(meta=meta, blocks=parse_blocks(text[match.end() :]))


def parse_blocks(body: str) -> list[Block]:
    blocks: list[Block] = []
    for chunk in re.split(r"\n\s*\n", body.strip()):
        lines = [line.rstrip() for line in chunk.strip().splitlines() if line.strip()]
        if not lines:
            continue
        first = lines[0]
        if first.startswith("## "):
            blocks.append(Block("heading", first[3:].strip()))
            rest = lines[1:]
            if rest:
                blocks.extend(parse_blocks("\n".join(rest)))
        elif first.startswith("!hidden "):
            blocks.append(Block("hidden", " ".join(lines)[len("!hidden ") :]))
        elif first.startswith("|"):
            rows = [
                [cell.strip() for cell in line.strip().strip("|").split("|")]
                for line in lines
                if not re.fullmatch(r"\|?[\s:|-]+\|?", line)
            ]
            blocks.append(Block("table", rows=rows))
        elif first.startswith("- "):
            blocks.append(Block("bullets", items=_list_items(lines, r"- ")))
        elif re.match(r"\d+\. ", first):
            blocks.append(Block("numbered", items=_list_items(lines, r"\d+\. ")))
        else:
            blocks.append(Block("paragraph", " ".join(lines)))
    return blocks


def _list_items(lines: list[str], marker: str) -> list[str]:
    items: list[str] = []
    for line in lines:
        if re.match(marker, line):
            items.append(re.sub(rf"^{marker}", "", line).strip())
        elif items:  # continuation line
            items[-1] += " " + line.strip()
    return items


def plain(text: str) -> str:
    return text.replace("**", "")


def human_date(raw: str) -> str:
    value = date.fromisoformat(raw)
    return f"{value.day} {value.strftime('%B')} {value.year}"


def title_lines(doc: SourceDocument) -> list[str]:
    """Metadata lines printed under the title. The ingestion pipeline reads
    them back (document ID, version, effective date, status)."""
    meta = doc.meta
    if doc.external:
        return [f"Issued by: {meta.get('issuer', 'external party')}"]
    return [
        f"Document ID: {meta['doc_id']} | Version: {meta['version']} | "
        f"Effective date: {human_date(meta['effective_date'])} | Status: {meta['status']}",
        f"Owner: {meta['owner']} | Classification: Internal",
    ]


def running_footer(doc: SourceDocument) -> str:
    if doc.external:
        return "Falcon Logistics SAS"
    return f"Nexora Industries | {doc.meta['doc_id']} v{doc.meta['version']} | Internal"


# --------------------------------------------------------------------------------- PDF


def find_font(name: str) -> Path:
    for folder in FONT_DIRS:
        if folder and (Path(folder) / name).exists():
            return Path(folder) / name
    sys.exit(f"Font {name} not found. Install DejaVu Sans or set NEXORA_FONT_DIR.")


class DocumentPDF(FPDF):
    def __init__(self, doc: SourceDocument) -> None:
        super().__init__(format="A4")
        self.doc = doc
        self.add_font("DejaVu", "", find_font("DejaVuSans.ttf"))
        self.add_font("DejaVu", "B", find_font("DejaVuSans-Bold.ttf"))
        self.set_margins(20, 20, 20)
        self.set_auto_page_break(auto=True, margin=22)
        self.set_creation_date(FIXED_TIMESTAMP)
        self.set_title(doc.meta["title"])
        self.set_author("Falcon Logistics SAS" if doc.external else "Nexora Industries")
        self.set_creator("build_sample_documents.py")

    def header(self) -> None:
        if self.page_no() == 1:
            return
        self.set_font("DejaVu", "", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, plain(self.doc.meta["title"]), align="L")
        self.ln(10)
        self.set_text_color(0, 0, 0)

    def footer(self) -> None:
        self.set_y(-15)
        self.set_font("DejaVu", "", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, f"{running_footer(self.doc)} | Page {self.page_no()} of {{nb}}", align="C")
        self.set_text_color(0, 0, 0)


def render_pdf(doc: SourceDocument, path: Path) -> None:
    pdf = DocumentPDF(doc)
    pdf.add_page()
    width = pdf.epw

    pdf.set_font("DejaVu", "B", 9)
    pdf.set_text_color(31, 78, 121)
    pdf.cell(width, 5, "FALCON LOGISTICS" if doc.external else "NEXORA INDUSTRIES")
    pdf.ln(7)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("DejaVu", "B", 17)
    pdf.multi_cell(width, 8, plain(doc.meta["title"]), **NEXT_LINE)
    pdf.ln(1)
    pdf.set_font("DejaVu", "", 8.5)
    pdf.set_text_color(80, 80, 80)
    for line in title_lines(doc):
        pdf.multi_cell(width, 4.6, line, **NEXT_LINE)
    pdf.multi_cell(width, 4.6, EXTERNAL_DISCLAIMER if doc.external else DISCLAIMER, **NEXT_LINE)
    pdf.set_text_color(0, 0, 0)
    pdf.ln(2)
    pdf.set_draw_color(31, 78, 121)
    pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + width, pdf.get_y())
    pdf.ln(4)

    for block in doc.blocks:
        if block.kind == "heading":
            pdf.ln(2)
            pdf.set_font("DejaVu", "B", 12)
            pdf.set_text_color(31, 78, 121)
            pdf.multi_cell(width, 6.5, block.text, **NEXT_LINE)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        elif block.kind == "paragraph":
            pdf.set_font("DejaVu", "", 10)
            pdf.multi_cell(width, 5.2, block.text, markdown=True, **NEXT_LINE)
            pdf.ln(2)
        elif block.kind == "hidden":
            # White text in a tiny font: invisible when the PDF is displayed,
            # but extracted like any other text.
            pdf.set_font("DejaVu", "", 5)
            pdf.set_text_color(255, 255, 255)
            pdf.multi_cell(width, 2.5, block.text, **NEXT_LINE)
            pdf.set_text_color(0, 0, 0)
        elif block.kind in ("bullets", "numbered"):
            pdf.set_font("DejaVu", "", 10)
            for number, item in enumerate(block.items, start=1):
                marker = "•" if block.kind == "bullets" else f"{number})"
                pdf.set_x(pdf.l_margin + 3)
                pdf.cell(7, 5.2, marker)
                pdf.multi_cell(width - 10, 5.2, item, markdown=True, **NEXT_LINE)
                pdf.ln(0.8)
            pdf.ln(1.5)
        elif block.kind == "table":
            pdf.set_font("DejaVu", "", 9)
            heading = FontFace(emphasis="BOLD", fill_color=(221, 231, 241))
            with pdf.table(
                width=width, line_height=5, headings_style=heading, padding=1.5, text_align="LEFT"
            ) as table:
                for cells in block.rows:
                    row = table.row()
                    for cell in cells:
                        row.cell(plain(cell))
            pdf.ln(3)
    pdf.output(str(path))


# -------------------------------------------------------------------------------- DOCX


def _add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, rendering **bold** spans."""
    for index, part in enumerate(text.split("**")):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def render_docx(doc: SourceDocument, path: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    section = document.sections[0]
    section.header.paragraphs[0].text = "Nexora Industries | Internal"
    section.footer.paragraphs[0].text = running_footer(doc)
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    company = document.add_paragraph()
    run = company.add_run("NEXORA INDUSTRIES")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(31, 78, 121)
    document.add_paragraph(plain(doc.meta["title"]), style="Title")
    for line in [*title_lines(doc), DISCLAIMER]:
        paragraph = document.add_paragraph()
        meta_run = paragraph.add_run(line)
        meta_run.font.size = Pt(8.5)
        meta_run.font.color.rgb = RGBColor(80, 80, 80)

    for block in doc.blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=1)
        elif block.kind == "paragraph":
            _add_runs(document.add_paragraph(), block.text)
        elif block.kind == "bullets":
            for item in block.items:
                _add_runs(document.add_paragraph(style="List Bullet"), item)
        elif block.kind == "numbered":
            for item in block.items:
                _add_runs(document.add_paragraph(style="List Number"), item)
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=len(block.rows[0]))
            table.style = "Light Grid Accent 1"
            for cells in block.rows:
                row = table.add_row().cells
                for target, cell in zip(row, cells, strict=True):
                    target.text = plain(cell)
            document.add_paragraph()
        elif block.kind == "hidden":
            raise ValueError("hidden text is only supported in PDF sources")

    properties = document.core_properties
    properties.title = plain(doc.meta["title"])
    properties.author = "Nexora Industries"
    properties.created = properties.modified = FIXED_TIMESTAMP.replace(tzinfo=None)
    properties.last_modified_by = "Nexora Industries"
    properties.revision = 1
    document.save(str(path))


# --------------------------------------------------------------------------------- TXT


def render_txt(doc: SourceDocument, path: Path) -> None:
    wrap = textwrap.TextWrapper(width=92, break_long_words=False, break_on_hyphens=False)
    lines = ["NEXORA INDUSTRIES", plain(doc.meta["title"]), *title_lines(doc)]
    lines += [DISCLAIMER, ""]
    for block in doc.blocks:
        if block.kind == "heading":
            lines += [block.text, ""]
        elif block.kind == "paragraph":
            lines += [*wrap.wrap(plain(block.text)), ""]
        elif block.kind in ("bullets", "numbered"):
            for number, item in enumerate(block.items, start=1):
                marker = "- " if block.kind == "bullets" else f"{number}) "
                item_wrap = textwrap.TextWrapper(
                    width=92, initial_indent=marker, subsequent_indent=" " * len(marker)
                )
                lines += item_wrap.wrap(plain(item))
            lines.append("")
        elif block.kind == "table":
            lines += [" | ".join(plain(cell) for cell in cells) for cells in block.rows]
            lines.append("")
        elif block.kind == "hidden":
            raise ValueError("hidden text is only supported in PDF sources")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "txt": render_txt}


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    sources = sorted(SOURCE_DIR.glob("*.md"))
    expected = set()
    for source_path in sources:
        doc = parse_source(source_path)
        fmt = doc.meta["format"]
        if fmt not in RENDERERS:
            raise ValueError(f"{source_path.name}: unknown format {fmt!r}")
        output = OUTPUT_DIR / doc.meta["file"]
        RENDERERS[fmt](doc, output)
        expected.add(output.name)
        print(f"{output.name:<70} {output.stat().st_size / 1024:6.1f} KB")
    stale = [path for path in OUTPUT_DIR.iterdir() if path.is_file() and path.name not in expected]
    for path in stale:
        print(f"warning: {path.name} has no source", file=sys.stderr)
    print(f"{len(expected)} documents written to {OUTPUT_DIR.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
